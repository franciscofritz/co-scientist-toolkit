"""
pptx_helpers — reusable python-pptx scaffolding for themed 16:9 decks.

Usage pattern (one deck per script):

    import pptx_helpers as D
    s = D.title_slide("My Talk", "A subtitle", "Author · 2026")
    D.notes(s, "[0:00] speaker script for this slide ...")

    s = D.slide(); D.header(s, "A content slide", kicker="SECTION", tag="P5", tagcolor=D.BLUE)
    D.body(s, [("a top bullet", 0, D.INK, True), ("a sub-bullet", 1)])
    D.pic(s, "fig/diagram.png", D.Inches(2.0), width=D.Inches(9.0))   # centred
    D.notes(s, "[~2 min] ...")

    D.divider("Part II", "the argument")

    D.add_footers()                      # page numbers on every slide but the title
    D.save("/path/to/Deck.pptx")

Key conventions baked in:
- 16:9 (13.333 x 7.5 in).
- Calibri, a fixed teal/ink palette (override the module constants if you like).
- Equations and diagrams go in as PNG images (see figures.py) — NOT as live text,
  because PowerPoint fonts lack ∑ ∫ 𝓘 etc. Avoid fancy Unicode (subscripts, script
  letters) in slide TEXT — use plain words, or render the symbol as an image.
- Put a timed speaker script in EVERY slide's notes(). It is what makes a deck a talk.
- Text boxes DO NOT CLIP. A figure placed over text collides silently. Before trusting
  a layout, use overlaps_ok() to check figure_bottom < next_text_top, then render the
  finished .pptx with LibreOffice and eyeball it (see the skill's Pipeline C).
"""
import struct
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

# ---- palette (override after import if a different theme is wanted) ----
INK=RGBColor(0x1A,0x24,0x33); PRIMARY=RGBColor(0x0F,0x5C,0x6B); PRIMARY_D=RGBColor(0x0B,0x3D,0x47)
BLUE=RGBColor(0x1F,0x6F,0xEB); GREEN=RGBColor(0x1B,0x7A,0x3D); AMBER=RGBColor(0xA6,0x68,0x0B)
RED=RGBColor(0x9A,0x2A,0x2A); LIGHT=RGBColor(0xEE,0xF3,0xF5); MUTED=RGBColor(0x6B,0x77,0x85)
WHITE=RGBColor(0xFF,0xFF,0xFF); PANEL=RGBColor(0xF6,0xF9,0xFA); ACCENT=RGBColor(0x2E,0x85,0x97)
FONT="Calibri"

prs=Presentation(); prs.slide_width=Inches(13.333); prs.slide_height=Inches(7.5)
SW,SH=prs.slide_width,prs.slide_height; BLANK=prs.slide_layouts[6]

def slide(): return prs.slides.add_slide(BLANK)

def rect(s,l,t,w,h,fill,line=None,lw=1.0):
    sp=s.shapes.add_shape(MSO_SHAPE.RECTANGLE,l,t,w,h); sp.shadow.inherit=False
    if fill is None: sp.fill.background()
    else: sp.fill.solid(); sp.fill.fore_color.rgb=fill
    if line is None: sp.line.fill.background()
    else: sp.line.color.rgb=line; sp.line.width=Pt(lw)
    return sp

def rrect(s,l,t,w,h,fill):
    sp=s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,l,t,w,h); sp.shadow.inherit=False
    try: sp.adjustments[0]=0.5
    except Exception: pass
    sp.fill.solid(); sp.fill.fore_color.rgb=fill; sp.line.fill.background(); return sp

def tb(s,l,t,w,h):
    b=s.shapes.add_textbox(l,t,w,h); tf=b.text_frame; tf.word_wrap=True
    for m in ("margin_left","margin_right"): setattr(tf,m,Inches(0.05))
    for m in ("margin_top","margin_bottom"): setattr(tf,m,Inches(0.02))
    return tf

def para(tf,text,size,color,bold=False,italic=False,align=PP_ALIGN.LEFT,first=False,before=4,after=4,line=1.05):
    p=tf.paragraphs[0] if first else tf.add_paragraph()
    p.alignment=align; p.space_before=Pt(before); p.space_after=Pt(after)
    try: p.line_spacing=line
    except Exception: pass
    r=p.add_run(); r.text=text; f=r.font
    f.size=Pt(size); f.bold=bold; f.italic=italic; f.name=FONT; f.color.rgb=color
    return p

def _indent(p,marL,ind):
    pPr=p._p.get_or_add_pPr(); pPr.set('marL',str(int(marL))); pPr.set('indent',str(int(ind)))

def header(s,title,kicker=None,tag=None,tagcolor=PRIMARY):
    """Top accent bar + title (+ optional small grey KICKER above, + optional coloured TAG chip top-right).
    Keep `title` to ONE line (~<=48 chars): longer titles wrap into the accent bar, or run UNDER the
    tag chip. The overlap pre-check won't catch this — only the LibreOffice render will. Keep it short."""
    rect(s,0,0,SW,Inches(0.16),PRIMARY)
    tt=tb(s,Inches(0.55),Inches(0.32),Inches(10.6),Inches(1.05))
    if kicker: para(tt,kicker,12.5,MUTED,bold=True,first=True,after=2); para(tt,title,26,PRIMARY,bold=True,before=0)
    else: para(tt,title,27,PRIMARY,bold=True,first=True)
    rect(s,Inches(0.57),Inches(1.40),Inches(1.5),Pt(3),PRIMARY)
    if tag:
        bx=rrect(s,SW-Inches(2.4)-Inches(0.5),Inches(0.42),Inches(2.4),Inches(0.5),tagcolor)
        t2=bx.text_frame; t2.word_wrap=False; para(t2,tag,13,WHITE,bold=True,align=PP_ALIGN.CENTER,first=True)

def body(s,items,l=Inches(0.6),t=Inches(1.7),w=Inches(12.1),size=17,gap=9):
    """items: list of (text, level[, color][, bold]). level 0 = ●, 1 = –, 2 = ·."""
    tf=tb(s,l,t,w,Inches(5.0)); first=True; glyph={0:"●",1:"–",2:"·"}
    for it in items:
        text,lvl=it[0],it[1]; color=it[2] if len(it)>2 and it[2] else INK; bold=it[3] if len(it)>3 else False
        sz=size if lvl==0 else size-1.5
        p=para(tf,f"{glyph.get(lvl,'·')}  {text}",sz,color,bold=bold,first=first,before=(gap if lvl==0 else 2),after=2)
        _indent(p,Inches(0.30)*(lvl+1),Inches(-0.26)); first=False
    return tf

def pic(s,path,top,height=None,width=None,center=True,left=None):
    """Place an image. Give EITHER width OR height; the other scales to the image's true aspect."""
    kw={}
    if height: kw['height']=height
    if width: kw['width']=width
    p=s.shapes.add_picture(path,Inches(0),top,**kw)
    if center: p.left=int((SW-p.width)/2)
    elif left is not None: p.left=left
    return p

def notes(s,txt): s.notes_slide.notes_text_frame.text=txt

def title_slide(title,subtitle="",byline="",extra=""):
    s=slide(); rect(s,0,0,SW,SH,PRIMARY_D); rect(s,0,Inches(2.5),SW,Inches(0.05),ACCENT)
    tf=tb(s,Inches(0.9),Inches(1.1),Inches(11.5),Inches(1.4)); para(tf,title,46,WHITE,bold=True,first=True)
    if subtitle:
        t2=tb(s,Inches(0.92),Inches(2.7),Inches(11.5),Inches(1.0)); para(t2,subtitle,26,RGBColor(0xBE,0xDB,0xE0),first=True)
    t3=tb(s,Inches(0.92),Inches(4.4),Inches(11.5),Inches(1.6))
    if byline: para(t3,byline,15,WHITE,first=True,after=12)
    if extra: para(t3,extra,13.5,RGBColor(0x9F,0xC2,0xC9))
    return s

def divider(title,sub=""):
    s=slide(); rect(s,0,0,SW,SH,PRIMARY_D); rect(s,0,Inches(3.35),SW,Inches(0.04),ACCENT)
    tf=tb(s,Inches(0.9),Inches(2.7),Inches(11.5),Inches(1.4)); para(tf,title,40,WHITE,bold=True,first=True)
    if sub:
        t2=tb(s,Inches(0.92),Inches(3.6),Inches(11.5),Inches(0.9)); para(t2,sub,18,RGBColor(0xBE,0xDB,0xE0),first=True)
    return s

def add_footers(label="", skip_first=True):
    for i,sl in enumerate(prs.slides):
        if skip_first and i==0: continue
        ft=tb(sl,Inches(0.55),SH-Inches(0.42),Inches(8),Inches(0.32)); para(ft,label,10,MUTED,first=True)
        pg=tb(sl,SW-Inches(1.4),SH-Inches(0.42),Inches(0.9),Inches(0.32)); para(pg,str(i+1),10,MUTED,align=PP_ALIGN.RIGHT,first=True)

def save(path): prs.save(path); return path

# ---- overlap pre-check: figures placed over text collide silently ----
def png_dims(path):
    """(width_px, height_px) from PNG header — no PIL needed."""
    with open(path,'rb') as f:
        head=f.read(24)
    if head[:8]!=b'\x89PNG\r\n\x1a\n': raise ValueError(f"{path} is not a PNG file; figures must be PNG (matplotlib savefig to .png)")
    return struct.unpack('>II',head[16:24])

def fig_bottom_in(path, top_in, width_in=None, height_in=None):
    """Where (in inches) the bottom of a placed image lands, using its TRUE tight-cropped aspect."""
    w,h=png_dims(path); asp=w/h
    ph = (width_in/asp) if width_in else height_in
    return top_in + ph

def overlaps_ok(path, top_in, width_in, next_text_top_in, margin=0.05):
    """True if a width-placed figure ends above the next text block (with a small margin)."""
    return fig_bottom_in(path, top_in, width_in=width_in) <= next_text_top_in - margin
